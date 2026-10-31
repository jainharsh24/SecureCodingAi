# -*- coding: utf-8 -*-
"""
Script to generate the SecureCode AI Research Paper in Word (.docx) format.
Follows the formatting, layout, structure, and academic rigor of researchpaper_sample.pdf
and draws all verified content directly from finalreport.pdf and the Literature Review dataset.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_document():
    doc = docx.Document()
    
    # Configure base style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(9.5)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style.paragraph_format.line_spacing = 1.12
    style.paragraph_format.space_after = Pt(3.5)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    return doc

def set_page_margins(section, top=0.75, bottom=0.75, left=0.75, right=0.75):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)

def set_two_columns(section, space_inches=0.35):
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')
    space_twips = str(int(space_inches * 1440))
    if cols:
        cols[0].set(qn('w:num'), '2')
        cols[0].set(qn('w:space'), space_twips)
    else:
        cols_elm = OxmlElement('w:cols')
        cols_elm.set(qn('w:num'), '2')
        cols_elm.set(qn('w:space'), space_twips)
        sectPr.append(cols_elm)

def set_single_column(section):
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), '1')
    else:
        cols_elm = OxmlElement('w:cols')
        cols_elm.set(qn('w:num'), '1')
        sectPr.append(cols_elm)

def add_header_footer(doc):
    for idx, section in enumerate(doc.sections):
        # Footer
        footer = section.footer
        footer.is_linked_to_previous = False
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_p.text = ""
        r_left = f_p.add_run("SecureCode AI: Hands-On Secure Coding Practice & Multi-Engine Evaluation Platform  |  ")
        r_left.font.name = 'Times New Roman'
        r_left.font.size = Pt(8.5)
        r_left.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        
        # Add Page Number field
        r_page = f_p.add_run()
        r_page.font.name = 'Times New Roman'
        r_page.font.size = Pt(8.5)
        r_page.font.bold = True
        r_page.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        r_page._r.append(fldChar1)
        r_page._r.append(instrText)
        r_page._r.append(fldChar2)
        r_page._r.append(fldChar3)

def add_title_block(doc):
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.paragraph_format.line_spacing = 1.15
    r_title = p_title.add_run("SecureCode AI: A Comprehensive Hands-On Secure Coding Practice and Multi-Engine Evaluation Platform")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    
    # Authors
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(2)
    p_auth.paragraph_format.space_after = Pt(4)
    authors = [
        ("Harsh Jain", "a"),
        ("Atharva Kadam", "b"),
        ("Sayli Vekhande", "c"),
        ("Arya Ratnakar", "d")
    ]
    for i, (name, sup) in enumerate(authors):
        r_n = p_auth.add_run(name)
        r_n.font.name = 'Times New Roman'
        r_n.font.size = Pt(10)
        r_n.font.bold = True
        r_s = p_auth.add_run(f" {sup}")
        r_s.font.name = 'Times New Roman'
        r_s.font.size = Pt(7)
        r_s.font.superscript = True
        if i < len(authors) - 1:
            r_sep = p_auth.add_run("   ,   ")
            r_sep.font.name = 'Times New Roman'
            r_sep.font.size = Pt(10)
            
    # Affiliation
    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff.paragraph_format.space_before = Pt(2)
    p_aff.paragraph_format.space_after = Pt(2)
    r_aff = p_aff.add_run("a, b, c, d Department of Computer Engineering, TPCT’s Terna Engineering College, Nerul, Navi Mumbai 400706, India")
    r_aff.font.name = 'Times New Roman'
    r_aff.font.size = Pt(9)
    r_aff.font.italic = True
    r_aff.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    
    # Emails
    p_em = doc.add_paragraph()
    p_em.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_em.paragraph_format.space_before = Pt(0)
    p_em.paragraph_format.space_after = Pt(3)
    r_em = p_em.add_run("harshjain2324@ternaengg.ac.in | atharvakadam2324@ternaengg.ac.in | saylivekhande2324@ternaengg.ac.in | aryaratnakar2324@ternaengg.ac.in")
    r_em.font.name = 'Times New Roman'
    r_em.font.size = Pt(8)
    r_em.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Guidance Note
    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_guide.paragraph_format.space_before = Pt(1)
    p_guide.paragraph_format.space_after = Pt(12)
    r_g = p_guide.add_run("Under the Guidance of: Prof. Mohini Misale (Department of Computer Engineering, Terna Engineering College)")
    r_g.font.name = 'Times New Roman'
    r_g.font.size = Pt(8.5)
    r_g.font.bold = True
    r_g.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x22, 0x55)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_p(doc, text, bold_prefix=None, space_after=3.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.left_indent = Inches(0.2)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    return p

def add_math(doc, formula, eq_num=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(formula)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    r.font.italic = True
    if eq_num:
        r_tab = p.add_run(f"    \t({eq_num})")
        r_tab.font.name = 'Times New Roman'
        r_tab.font.size = Pt(9.5)
        r_tab.font.bold = True
        r_tab.font.italic = False

def add_figure(doc, img_name, caption_text, width=Inches(3.3)):
    img_path = os.path.join('images_&_ppt', img_name)
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found.")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(6)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(8.5)
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def style_academic_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    borders_xml = (
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        r'  <w:left w:val="none"/>'
        r'  <w:right w:val="none"/>'
        r'  <w:insideV w:val="none"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(parse_xml(borders_xml))
    
    # Header format
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    
    for c_idx, cell in enumerate(table.rows[0].cells):
        shd_xml = r'<w:shd %s w:fill="F2F4F8"/>' % nsdecls('w')
        cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(8.5)
                r.font.bold = True
                
    # Body rows formatting
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(8)
                    
    # Column widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

print("Base setup script defined.")
