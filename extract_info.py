import openpyxl
import json

wb = openpyxl.load_workbook('images_&_ppt/Literature_review(1-27 papers).xlsx')
sheet = wb.active
papers = []
for r in range(3, sheet.max_row + 1):
    sno = sheet.cell(r, 1).value
    if not sno:
        continue
    paper = {
        'sno': sno,
        'title': sheet.cell(r, 2).value,
        'authors': sheet.cell(r, 3).value,
        'journal': sheet.cell(r, 4).value,
        'year': sheet.cell(r, 5).value,
        'keywords': sheet.cell(r, 6).value,
        'objective': sheet.cell(r, 7).value,
        'type': sheet.cell(r, 8).value,
        'solution': sheet.cell(r, 9).value,
        'methodology': sheet.cell(r, 13).value,
        'findings': sheet.cell(r, 14).value,
        'limitations': sheet.cell(r, 15).value,
        'takeaways': sheet.cell(r, 16).value,
        'citation': sheet.cell(r, 17).value
    }
    papers.append(paper)

print(f'Extracted {len(papers)} papers.')
with open('extracted_papers.json', 'w', encoding='utf-8') as f:
    json.dump(papers, f, indent=2, ensure_ascii=False)

for p in papers:
    print(f"[{p['sno']}] {p['title']} ({p['year']}) - {p['authors']}")
