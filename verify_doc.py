import docx

doc = docx.Document('Report.docx')

print(f"Total Sections: {len(doc.sections)}")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# Check embedded images (shapes / drawings)
img_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.target_ref:
        img_count += 1
print(f"Total Embedded Image Parts: {img_count}")

# Check headings
headings = [p.text for p in doc.paragraphs if p.text.startswith("Chapter ") or p.text in ["Abstract", "References", "Declaration", "CERTIFICATE"]]
print(f"Found Major Headings: {headings}")

# Check tables
table_titles = [p.text for p in doc.paragraphs if p.text.startswith("Table ")]
print(f"Found Table Titles: {len(table_titles)}")
for tt in table_titles:
    print(f"  {tt}")

# Check figures
fig_captions = [p.text for p in doc.paragraphs if p.text.startswith("Fig ")]
print(f"Found Figure Captions: {len(fig_captions)}")
for fc in fig_captions:
    print(f"  {fc}")
