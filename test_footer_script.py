import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()
section = doc.sections[0]
footer = section.footer
f_p = footer.paragraphs[0]
f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

run = f_p.add_run()
fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
run._r.append(fldChar1)
run._r.append(instrText)
run._r.append(fldChar2)
run._r.append(fldChar3)
f_p.add_run(' | P a g e')

table = doc.add_table(rows=2, cols=2)
shading_xml = parse_xml(r'<w:shd %s w:fill="1F4E79"/>' % nsdecls('w'))
table.rows[0].cells[0]._tc.get_or_add_tcPr().append(shading_xml)

doc.save('test_footer.docx')
print('Footer and table test passed!')
