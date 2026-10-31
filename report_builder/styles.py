import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

NAVY_BLUE = RGBColor(31, 78, 121)     # Primary accent #1F4E79
DARK_GRAY = RGBColor(51, 51, 51)      # Text dark #333333
SLATE_GRAY = RGBColor(100, 116, 139)  # Subdued #64748B
BLACK = RGBColor(0, 0, 0)

HEX_NAVY = "1F4E79"
HEX_LIGHT_BG = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_ZEBRA = "F8FAFC"

def init_document():
    doc = docx.Document()
    # Configure 1 inch top/bottom/right, 1.25 inch left (binding margin)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        
        # Configure footer with dynamic page number
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_p.paragraph_format.space_before = Pt(6)
        f_p.paragraph_format.space_after = Pt(0)
        
        # Dynamic Page Field
        run = f_p.add_run()
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9.5)
        run.font.color.rgb = SLATE_GRAY
        
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run._r.append(fldChar3)
        
        run_label = f_p.add_run(" | P a g e")
        run_label.font.name = 'Times New Roman'
        run_label.font.size = Pt(9.5)
        run_label.font.color.rgb = SLATE_GRAY

    # Configure normal style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = DARK_GRAY
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    return doc

def add_p(doc, text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6, font_size=12, color=DARK_GRAY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
    return p

def add_bullet(doc, text, bold_prefix="", level=0, space_after=3):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11.5)
        r_pre.font.color.rgb = DARK_GRAY
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11.5)
    run.font.color.rgb = DARK_GRAY
    return p

def add_chapter_heading(doc, chapter_num, chapter_title):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    run_ch = p.add_run(f"Chapter {chapter_num}\n")
    run_ch.bold = True
    run_ch.font.name = 'Times New Roman'
    run_ch.font.size = Pt(18)
    run_ch.font.color.rgb = NAVY_BLUE
    
    run_title = p.add_run(chapter_title)
    run_title.bold = True
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = NAVY_BLUE
    
    # Subtle separator line or space
    p.paragraph_format.space_after = Pt(14)
    return p

def add_sec_heading(doc, sec_num, sec_title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{sec_num} {sec_title}")
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.color.rgb = NAVY_BLUE
    return p

def add_subsec_heading(doc, subsec_num, subsec_title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{subsec_num} {subsec_title}")
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12.5)
    run.font.color.rgb = DARK_GRAY
    return p

def add_subsubsec_heading(doc, subsubsec_title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(subsubsec_title)
    run.bold = True
    run.italic = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = DARK_GRAY
    return p

def add_figure(doc, img_path, fig_id, caption_text, width_inches=6.0):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    
    try:
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))
    except Exception as e:
        run_err = p_img.add_run(f"[Image could not be rendered: {img_path}]")
        run_err.italic = True
        
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    
    run_id = p_cap.add_run(f"{fig_id} ")
    run_id.bold = True
    run_id.font.name = 'Times New Roman'
    run_id.font.size = Pt(10.5)
    run_id.font.color.rgb = NAVY_BLUE
    
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(10.5)
    run_cap.italic = True
    run_cap.font.color.rgb = DARK_GRAY
    return p_cap

def create_styled_table(doc, headers, data, col_widths=None, table_title=None, table_id=None):
    if table_title and table_id:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(4)
        p_title.paragraph_format.keep_with_next = True
        
        run_id = p_title.add_run(f"{table_id}: ")
        run_id.bold = True
        run_id.font.name = 'Times New Roman'
        run_id.font.size = Pt(10.5)
        run_id.font.color.rgb = NAVY_BLUE
        
        run_title = p_title.add_run(table_title)
        run_title.bold = True
        run_title.font.name = 'Times New Roman'
        run_title.font.size = Pt(10.5)
        run_title.font.color.rgb = DARK_GRAY

    rows_cnt = len(data) + 1
    cols_cnt = len(headers)
    table = doc.add_table(rows=rows_cnt, cols=cols_cnt)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Apply borders and formatting via XML
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="%s"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="%s"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="%s"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''' % (nsdecls('w'), HEX_BORDER, HEX_NAVY, HEX_BORDER, HEX_BORDER))
    tblPr.append(borders_xml)

    # Style Header Row
    hdr_row = table.rows[0]
    # Set cantSplit and tblHeader
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))

    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), HEX_NAVY))
        cell._tc.get_or_add_tcPr().append(shd)
        
        # cell margins (padding)
        tcMar = parse_xml(r'''
            <w:tcMar %s>
                <w:top w:w="120" w:type="dxa"/>
                <w:bottom w:w="120" w:type="dxa"/>
                <w:left w:w="140" w:type="dxa"/>
                <w:right w:w="140" w:type="dxa"/>
            </w:tcMar>
        ''' % nsdecls('w'))
        cell._tc.get_or_add_tcPr().append(tcMar)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(str(h_text))
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Style Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        
        is_zebra = (r_idx % 2 == 1)
        bg_color = HEX_ZEBRA if is_zebra else "FFFFFF"

        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), bg_color))
            cell._tc.get_or_add_tcPr().append(shd)

            tcMar = parse_xml(r'''
                <w:tcMar %s>
                    <w:top w:w="100" w:type="dxa"/>
                    <w:bottom w:w="100" w:type="dxa"/>
                    <w:left w:w="140" w:type="dxa"/>
                    <w:right w:w="140" w:type="dxa"/>
                </w:tcMar>
            ''' % nsdecls('w'))
            cell._tc.get_or_add_tcPr().append(tcMar)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.10
            run = p.add_run(str(val) if val is not None else "")
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9.5)
            run.font.color.rgb = DARK_GRAY

    # Apply column widths if provided
    if col_widths and len(col_widths) == cols_cnt:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(2)
    p_post.paragraph_format.space_after = Pt(6)
    return table
