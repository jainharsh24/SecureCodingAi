from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_builder.styles import add_p, create_styled_table, NAVY_BLUE, DARK_GRAY, BLACK

def generate_frontmatter(doc):
    # ================= PAGE 1: TITLE PAGE =================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    run_t = p_title.add_run("SecureCode AI: A Comprehensive Hands-On Secure Coding Practice and Multi-Engine Evaluation Platform")
    run_t.bold = True
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(20)
    run_t.font.color.rgb = NAVY_BLUE

    add_p(doc, "Major Project I Report", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=15, space_before=10, space_after=4, color=DARK_GRAY)
    add_p(doc, "Submitted in partial fulfillment of the requirement of University of Mumbai", bold=False, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=2, space_after=2)
    add_p(doc, "For the Degree of", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=2, space_after=2)
    add_p(doc, "Bachelor of Engineering (Computer Engineering)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=14, space_before=2, space_after=14, color=NAVY_BLUE)

    add_p(doc, "By", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=6, space_after=6)
    
    students = [
        "1) Harsh Jain (ID No: TU3F2324058 / Roll No: A-58)",
        "2) Atharva Kadam (ID No: TU3F2324054 / Roll No: A-54)",
        "3) Sayli Vekhande (ID No: TU3F2324063 / Roll No: A-63)",
        "4) Arya Ratnaker (ID No: TU3F2324047 / Roll No: A-47)"
    ]
    for s in students:
        add_p(doc, s, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=2, space_after=2)

    add_p(doc, "Under the Guidance of", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=14, space_after=4)
    add_p(doc, "Prof. Mohini Misale", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=13, space_before=2, space_after=20, color=NAVY_BLUE)

    add_p(doc, "Department of Computer Engineering", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=13, space_before=10, space_after=2)
    add_p(doc, "TERNA ENGINEERING COLLEGE", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=14, space_before=2, space_after=2, color=NAVY_BLUE)
    add_p(doc, "Plot no. 12, Sector-22, Opp. Nerul Railway Station, Phase-II, Nerul (W), Navi Mumbai 400706", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11, space_before=2, space_after=2)
    add_p(doc, "UNIVERSITY OF MUMBAI", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=13, space_before=4, space_after=2)
    add_p(doc, "Academic Year 2026-27", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=2, space_after=0)

    # ================= PAGE 2: CERTIFICATE =================
    doc.add_page_break()
    add_p(doc, "TERNA ENGINEERING COLLEGE, NERUL, NAVI MUMBAI", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=14, space_before=16, space_after=2, color=NAVY_BLUE)
    add_p(doc, "Department of Computer Engineering", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=13, space_before=2, space_after=2)
    add_p(doc, "Academic Year 2026-27", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=2, space_after=24)

    add_p(doc, "CERTIFICATE", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=8, space_after=18, color=NAVY_BLUE)

    cert_text = (
        "This is to certify that the Major Project I entitled “SecureCode AI: A Comprehensive Hands-On "
        "Secure Coding Practice and Multi-Engine Evaluation Platform” is a bonafide work of:"
    )
    add_p(doc, cert_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=4, space_after=10)

    for s in students:
        add_p(doc, s, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12, space_before=2, space_after=2)

    cert_sub = (
        "submitted to the University of Mumbai in partial fulfillment of the requirement for the award "
        "of the Degree of Bachelor of Engineering (Computer Engineering)."
    )
    add_p(doc, cert_sub, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=10, space_after=60)

    # Signatures table
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_cells = sig_table.rows[0].cells
    sig_cells[0].paragraphs[0].text = "____________________\nProf. Mohini Misale\nGuide"
    sig_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_cells[1].paragraphs[0].text = "____________________\nDr. Kishor Sakure\nHead of Department"
    sig_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_cells[2].paragraphs[0].text = "____________________\nDr. Deven Shah\nPrincipal"
    sig_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ================= PAGE 3: APPROVAL =================
    doc.add_page_break()
    add_p(doc, "Project Report Approval", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=20, space_after=18, color=NAVY_BLUE)

    app_text = (
        "This Major Project I Report – entitled “SecureCode AI: A Comprehensive Hands-On Secure Coding "
        "Practice and Multi-Engine Evaluation Platform” by the following students is approved for the "
        "degree of B.E. in “Computer Engineering”."
    )
    add_p(doc, app_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=4, space_after=14)

    add_p(doc, "Submitted by:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12, space_before=6, space_after=6)
    for s in students:
        add_p(doc, s, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12, space_before=2, space_after=2)

    add_p(doc, "Examiners Name & Signature:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12, space_before=40, space_after=16)
    add_p(doc, "1. ____________________________________________________", font_size=12, space_before=8, space_after=14)
    add_p(doc, "2. ____________________________________________________", font_size=12, space_before=8, space_after=24)

    add_p(doc, "Date: _________________________", font_size=12, space_before=14, space_after=6)
    add_p(doc, "Place: ________________________", font_size=12, space_before=4, space_after=6)

    # ================= PAGE 4: DECLARATION =================
    doc.add_page_break()
    add_p(doc, "Declaration", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=20, space_after=18, color=NAVY_BLUE)

    decl_text = (
        "We declare that this written submission represents our ideas in our own words and where "
        "others' ideas or words have been included, we have adequately cited and referenced the "
        "original sources. We also declare that we have adhered to all principles of academic honesty "
        "and integrity and have not misrepresented or fabricated or falsified any idea/data/fact/source "
        "in our submission. We understand that any violation of the above will be cause for disciplinary "
        "action by the Institute and can also evoke penal action from the sources which have thus not been "
        "properly cited or from whom proper permission has not been taken when needed."
    )
    add_p(doc, decl_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=6, space_after=30)

    for s in students:
        add_p(doc, f"{s}    _________________________", bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11, space_before=6, space_after=6)

    add_p(doc, "Date: _________________________", font_size=12, space_before=24, space_after=6)
    add_p(doc, "Place: ________________________", font_size=12, space_before=4, space_after=6)

    # ================= PAGE 5: ACKNOWLEDGEMENT =================
    doc.add_page_break()
    add_p(doc, "Acknowledgement", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=20, space_after=18, color=NAVY_BLUE)

    ack1 = (
        "We would like to express our sincere gratitude towards our guide Prof. Mohini Misale and our "
        "Major Project Coordinator Dr. Siddharth Hariharan for their invaluable guidance, encouragement, "
        "and constant constructive feedback throughout the inception, design, and implementation of SecureCode AI. "
        "This project would not have been possible without their continuous support, scholarly direction, "
        "and motivation. We thank them for making our engineering endeavor thoroughly engaging and intellectually enriching."
    )
    add_p(doc, ack1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=6, space_after=10)

    ack2 = (
        "We are deeply thankful to Dr. Kishor Sakure (H.O.D., Department of Computer Engineering) and the faculty "
        "members of the Computer Engineering Department for providing access to computing facilities, software "
        "resources, and technical consultation. Their scientific rigor and pedagogical standards inspired us to strive "
        "for academic excellence and robust software engineering practices."
    )
    add_p(doc, ack2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=6, space_after=10)

    ack3 = (
        "We take the privilege to express our sincere thanks to Dr. Deven Shah, our Principal, for providing an "
        "environment that fosters technological innovation and research. Finally, we thank our peers and families "
        "for their steadfast encouragement throughout this endeavor."
    )
    add_p(doc, ack3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=6, space_after=24)

    for s in students:
        add_p(doc, f"{s}    _________________________", bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11, space_before=4, space_after=4)

    # ================= PAGE 6-7: TABLE OF CONTENTS =================
    doc.add_page_break()
    add_p(doc, "TABLE OF CONTENTS", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=16, space_after=12, color=NAVY_BLUE)

    toc_data = [
        ["01", "Abstract", "08"],
        ["02", "List of Figures", "09"],
        ["03", "List of Tables", "10"],
        ["Chapter 1", "Introduction", "11-14"],
        ["", "1.1 Motivation", "11"],
        ["", "    1.1.1 Need of the Problem", "12"],
        ["", "1.2 Scope of the Project", "13"],
        ["", "1.3 Aim and Research Objectives", "14"],
        ["Chapter 2", "Problem Statement & Requirements Engineering", "15-19"],
        ["", "2.1 Problem Statement", "15"],
        ["", "2.2 Features of the System", "16"],
        ["", "2.3 Objectives", "17"],
        ["", "2.4 System Specifications & Requirements", "17"],
        ["", "    2.4.1 Functional Requirements", "17"],
        ["", "    2.4.2 Non-Functional Requirements", "18"],
        ["Chapter 3", "Literature Review", "20-27"],
        ["", "3.1 Existing Systems", "20"],
        ["", "3.2 Limitations of Existing Systems", "22"],
        ["", "3.3 Proposed Solution", "23"],
        ["", "3.4 Research Gap & Comparative Analysis", "24"],
        ["", "3.5 Literature Survey (27 Peer-Reviewed Works)", "25"],
        ["Chapter 4", "Design and Implementation", "28-44"],
        ["", "4.1 System Architecture", "28"],
        ["", "4.2 Working of the System", "30"],
        ["", "4.3 Use Case Modeling", "31"],
        ["", "4.4 Data Flow Diagrams (Level 0, 1, and 2)", "32"],
        ["", "4.5 Evaluation Workflow & Activity Modeling", "35"],
        ["", "4.6 Multi-Engine Security Architecture & Consensus", "36"],
        ["", "4.7 Algorithms and Mathematical Formulations", "37"],
        ["", "    4.7.1 Dual-Stage Evaluation Gate Algorithm", "37"],
        ["", "    4.7.2 Static Analysis Normalization Algorithm", "38"],
        ["", "    4.7.3 Four-Pillar Overall Learning Score Formulation", "38"],
        ["", "    4.7.4 Attempt Windowing & Prerequisite Progression", "39"],
        ["", "4.8 Relational Data Model & Database Schema", "40"],
        ["", "4.9 Class Diagram & Domain Layer Services", "41"],
        ["", "4.10 UI Wireframe Design & Interaction Principles", "41"],
        ["", "4.11 Software Process Model & Project Timeline", "42"],
        ["", "4.12 Risk Mitigation, Monitoring and Management (RMMM) Plan", "43"],
        ["", "4.13 Hardware Requirements", "44"],
        ["", "4.14 Software Requirements & Technology Stack", "44"],
        ["Chapter 5", "Results and Discussions", "45-56"],
        ["", "5.1 Experimental Evaluation and Benchmarks", "45"],
        ["", "5.2 Landing & Home Page", "47"],
        ["", "5.3 Challenge Catalog & Vulnerability Selection", "48"],
        ["", "5.4 Learning Hub & Pedagogical Fundamentals", "49"],
        ["", "5.5 Interactive Code Editor & Dual-Stage Evaluation", "50"],
        ["", "5.6 Attempt History & Remediation Progression", "51"],
        ["", "5.7 Submission Analysis & Multi-Layer Feedback", "52"],
        ["", "5.8 Analytics Dashboard & Security Metrics", "53"],
        ["", "5.9 Vulnerability Trend Analysis & Injection Prioritization", "54"],
        ["", "5.10 Industry Platform Comparison & Gap Resolution", "55"],
        ["Chapter 6", "Conclusion and Future Scope", "57-58"],
        ["", "6.1 Conclusion", "57"],
        ["", "6.2 Future Scope", "58"],
        ["04", "References (Consolidated Bibliography)", "59-62"],
        ["05", "Appendix A: Active Injection Catalog & Rubric", "63-65"]
    ]

    create_styled_table(doc, ["Sr. No.", "Title", "Page No."], toc_data, col_widths=[1.2, 4.3, 0.9])

    # ================= PAGE 8: ABSTRACT =================
    doc.add_page_break()
    add_p(doc, "Abstract", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=16, space_after=14, color=NAVY_BLUE)

    ab1 = (
        "Software security vulnerabilities remain a prevalent and costly threat to global digital infrastructure. "
        "While traditional competitive programming and educational coding platforms (such as LeetCode, HackerRank, "
        "and Codewars) effectively cultivate algorithmic proficiency, they focus exclusively on functional correctness "
        "while ignoring software security vulnerabilities. Conversely, offensive cybersecurity platforms (such as PortSwigger "
        "Web Security Academy, TryHackMe, and Hack The Box) emphasize penetration testing and exploitation rather than the "
        "constructive software engineering disciplines required to write secure source code. Furthermore, modern generative "
        "AI coding assistants frequently produce functionally acceptable yet insecure code, exacerbating the vulnerability epidemic."
    )
    add_p(doc, ab1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=4, space_after=8)

    ab2 = (
        "This report presents SecureCode AI, an interactive, web-based secure coding practice and automated assessment "
        "platform engineered to cultivate defensive programming competencies from the ground up. SecureCode AI implements "
        "a novel dual-stage evaluation architecture: Stage 1 executes dynamic functional unit testing within an isolated sandbox "
        "(via Judge0), while Stage 2 performs multi-engine static and AI-assisted security analysis combining abstract syntax tree (AST) "
        "rules (Semgrep), bytecode dataflow inspection (SpotBugs/FindSecBugs), and contextual large language model advisory (Groq LLaMA-3). "
        "To prevent trial-and-error guessing and superficial code modification, SecureCode AI introduces an active catalog of 45 "
        "staged injection challenges (spanning SQL, NoSQL, OS Command, LDAP, XPath, Header/CRLF, and Code Injection), a rate-limiting "
        "five-attempt window, interactive conceptual understanding quizzes, and a multi-factor learning scoring engine that holistically "
        "weights Functional Correctness (25%), Security Correctness (45%), Conceptual Understanding (15%), and Independent Problem Solving (15%)."
    )
    add_p(doc, ab2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=4, space_after=8)

    ab3 = (
        "The system was engineered using Spring Boot 3, Java 25, PostgreSQL, Flyway migrations, and a responsive web application client. "
        "Empirical validation across 80 automated unit, integration, and security test suites demonstrates zero false passes on vulnerable code, "
        "flawless consensus arbitration across disparate static analysis outputs, and sub-500ms median evaluation latency for local static engines. "
        "SecureCode AI bridges the critical pedagogical chasm between algorithmic programming and software security, empowering developers "
        "to construct verifiably secure systems before deployment."
    )
    add_p(doc, ab3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12, space_before=4, space_after=14)

    add_p(doc, "Keywords: Secure Code Generation, Automated Programming Assessment Systems (APAS), Static Application Security Testing (SAST), Injection Attacks, Dual-Stage Evaluation, Learning Analytics, Large Language Models (LLM), Software Engineering Education.", 
          bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=11, space_before=6, space_after=6)

    # ================= PAGE 9: LIST OF FIGURES =================
    doc.add_page_break()
    add_p(doc, "List of Figures", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=16, space_after=12, color=NAVY_BLUE)

    fig_data = [
        ["Chapter 4", "", ""],
        ["Fig 4.1.1", "SecureCode AI Proposed Multi-Tier System Architecture", "29"],
        ["Fig 4.3.1", "Comprehensive Use Case Diagram for Students, Administrators, and Evaluation Engine", "31"],
        ["Fig 4.4.1", "Data Flow Diagram – Level 0 (Context Diagram)", "32"],
        ["Fig 4.4.2", "Data Flow Diagram – Level 1 (Core Subsystems and Data Stores)", "33"],
        ["Fig 4.4.3", "Data Flow Diagram – Level 2 (Evaluation Engine Decomposition & Normalization Pipeline)", "34"],
        ["Fig 4.5.1", "Activity Diagram & Submission Evaluation Workflow", "35"],
        ["Fig 4.11.1", "Gantt Chart (SecureCode AI 16-Week Implementation Timeline)", "42"],
        ["Chapter 5", "", ""],
        ["Fig 5.2.1", "SecureCode AI Landing Page and Pedagogical Overview", "47"],
        ["Fig 5.3.1", "Interactive Challenge Catalog with Tiered Difficulty and Search Filters", "48"],
        ["Fig 5.4.1", "Learning Hub with Vulnerability Classifications and Secure Principles", "49"],
        ["Fig 5.5.1", "Browser-Based Code Editor, Test Suite Runner, and Multi-Tiered Hints", "50"],
        ["Fig 5.6.1", "Five-Attempt Remediation History & Iterative Learning Score Progression", "51"],
        ["Fig 5.7.1", "Submission Analysis – Code Inspection, Multi-Scanner Consensus, and AI Advisory", "52"],
        ["Fig 5.7.2", "Submission Analysis – Four-Pillar Learning Score Breakdown & Conceptual Assessment", "53"],
        ["Fig 5.8.1", "Student Analytics Dashboard – Progress Metrics & Evaluation Accuracy", "54"],
        ["Fig 5.9.1", "Vulnerability Classification Tree and Security Research Landscape", "55"],
        ["Fig 5.10.1", "Comparative Analysis of Existing Industry Platforms and Identification of Research Gap", "56"]
    ]
    create_styled_table(doc, ["Figure No.", "Name of Figure", "Page No."], fig_data, col_widths=[1.3, 4.3, 0.8])

    # ================= PAGE 10: LIST OF TABLES =================
    doc.add_page_break()
    add_p(doc, "List of Tables", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16, space_before=16, space_after=12, color=NAVY_BLUE)

    tbl_data = [
        ["Chapter 2", "", ""],
        ["Table 2.4.1", "Comprehensive Non-Functional Requirements Specification", "18"],
        ["Chapter 3", "", ""],
        ["Table 3.5.1", "Literature Survey Summary of 27 Peer-Reviewed Research Papers", "25"],
        ["Chapter 4", "", ""],
        ["Table 4.8.1", "Relational Database Schema and Entity Descriptions (13 Core Tables)", "40"],
        ["Table 4.11.1", "Comparative Evaluation and Selection of the Software Engineering Model", "42"],
        ["Table 4.11.2", "Iterative Development Milestones, Features Delivered, and Review Status", "43"],
        ["Table 4.12.1", "Risk Mitigation, Monitoring, and Management (RMMM) Plan", "43"],
        ["Table 4.14.1", "Tools, Frameworks, Libraries, and Technologies Used in SecureCode AI", "44"],
        ["Chapter 5", "", ""],
        ["Table 5.1.1", "Distribution of Active Injection Catalog by Vulnerability Subtype and Difficulty", "45"],
        ["Table 5.1.2", "Dual-Stage Evaluation Performance, Execution Latency, and Test Benchmarks", "46"],
        ["Table 5.1.3", "Four-Pillar Learning Score Mathematical Weighting and Scoring Matrix", "46"],
        ["Table 5.10.1", "Detailed Feature-by-Feature Matrix Comparing SecureCode AI with Industry Platforms", "56"],
        ["Appendix A", "", ""],
        ["Table A.1", "Active Injection Challenge Catalog Benchmark Suite and CWE Mappings", "63"]
    ]
    create_styled_table(doc, ["Table No.", "Name of Table", "Page No."], tbl_data, col_widths=[1.3, 4.3, 0.8])
