import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Theme Colors
NAVY = RGBColor(31, 78, 121)        # #1F4E79
DARK = RGBColor(34, 34, 34)         # #222222
SLATE = RGBColor(100, 116, 139)     # #64748B
RED = RGBColor(192, 0, 0)           # #C00000
LIGHT_BG = RGBColor(241, 245, 249)  # #F1F5F9
WHITE = RGBColor(255, 255, 255)

FONT_NAME = "Times New Roman"
LOGO_PATH = "images_&_ppt/college_logo.png"

def add_header(slide, title_text, slide_num, total_slides=35):
    """Adds college logo, standard title header, and slide counter."""
    # College Logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(0.35), Inches(0.22), Inches(1.80), Inches(0.80))
    
    # Title Box
    title_box = slide.shapes.add_textbox(Inches(2.30), Inches(0.18), Inches(9.80), Inches(0.88))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_NAME
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.LEFT
    
    # Slide Number
    num_box = slide.shapes.add_textbox(Inches(11.80), Inches(7.05), Inches(1.20), Inches(0.35))
    ntf = num_box.text_frame
    ntf.word_wrap = False
    np = ntf.paragraphs[0]
    np.text = f"{slide_num} / {total_slides}"
    np.font.name = FONT_NAME
    np.font.size = Pt(12)
    np.font.color.rgb = SLATE
    np.alignment = PP_ALIGN.RIGHT

def add_bullet_content(slide, bullets, top_inch=1.35, height_inch=5.50, width_inch=12.40, left_inch=0.55, font_size=20):
    """Adds clean, professional bullet points."""
    content_box = slide.shapes.add_textbox(Inches(left_inch), Inches(top_inch), Inches(width_inch), Inches(height_inch))
    tf = content_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.font.name = FONT_NAME
        p.font.size = Pt(font_size)
        p.font.color.rgb = DARK
        p.space_after = Pt(12)
        p.space_before = Pt(4)
        p.level = 0
        
        # If bullet has bold prefix (e.g. "Prefix: text")
        if ":" in b and not b.startswith("http"):
            parts = b.split(":", 1)
            r_pre = p.add_run()
            r_pre.text = parts[0] + ":"
            r_pre.font.bold = True
            r_pre.font.name = FONT_NAME
            r_pre.font.size = Pt(font_size)
            r_pre.font.color.rgb = NAVY
            
            r_body = p.add_run()
            r_body.text = parts[1]
            r_body.font.bold = False
            r_body.font.name = FONT_NAME
            r_body.font.size = Pt(font_size)
            r_body.font.color.rgb = DARK
        else:
            r = p.add_run()
            r.text = b
            r.font.name = FONT_NAME
            r.font.size = Pt(font_size)
            r.font.color.rgb = DARK

def add_image_slide(slide, img_path, takeaway_bullets=None, max_img_w=12.2, max_img_h=4.8, img_top=1.30):
    """Centers an image and optionally places a takeaway banner/bullets at bottom."""
    from PIL import Image
    im = Image.open(img_path)
    orig_w, orig_h = im.size
    aspect = orig_w / orig_h
    
    # Calculate fit
    w = max_img_w
    h = w / aspect
    if h > max_img_h:
        h = max_img_h
        w = h * aspect
    
    left = (13.333 - w) / 2
    top = img_top
    
    slide.shapes.add_picture(img_path, Inches(left), Inches(top), Inches(w), Inches(h))
    
    if takeaway_bullets:
        box_top = top + h + 0.15
        box_h = 7.10 - box_top
        c_box = slide.shapes.add_textbox(Inches(0.65), Inches(box_top), Inches(12.0), Inches(box_h))
        tf = c_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        for i, tb in enumerate(takeaway_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.font.name = FONT_NAME
            p.font.size = Pt(17)
            p.space_after = Pt(4)
            p.space_before = Pt(2)
            if ":" in tb:
                parts = tb.split(":", 1)
                r_pre = p.add_run()
                r_pre.text = parts[0] + ":"
                r_pre.font.bold = True
                r_pre.font.name = FONT_NAME
                r_pre.font.color.rgb = NAVY
                r_body = p.add_run()
                r_body.text = parts[1]
                r_body.font.name = FONT_NAME
                r_body.font.color.rgb = DARK
            else:
                r = p.add_run()
                r.text = tb
                r.font.name = FONT_NAME
                r.font.color.rgb = DARK

def add_table_slide(slide, headers, rows, col_widths, top_inch=1.40):
    """Creates a beautifully styled, high-contrast academic table."""
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(0.55), Inches(top_inch), Inches(12.20), Inches(5.30))
    table = table_shape.table
    
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
        
    # Header Row
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_NAME
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
    # Data Rows
    for row_idx, r_data in enumerate(rows):
        is_even = (row_idx % 2 == 1)
        bg = LIGHT_BG if is_even else WHITE
        for col_idx, val in enumerate(r_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.name = FONT_NAME
            p.font.size = Pt(13)
            p.font.color.rgb = DARK
            p.alignment = PP_ALIGN.CENTER if col_idx in [0, 1] else PP_ALIGN.LEFT
